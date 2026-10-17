import json
import docx
import pptx
from bs4 import BeautifulSoup
from pypdf import PdfReader
from langchain_text_splitters import RecursiveCharacterTextSplitter

text_splitter = RecursiveCharacterTextSplitter(chunk_size=1200, chunk_overlap=200)

def extract_docx(file_path):
    doc = docx.Document(file_path)
    units = []
    current_heading = "Overview"
    current_paras = []
    
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        if p.style and p.style.name and p.style.name.startswith("Heading"):
            if current_paras:
                units.append({"section": current_heading, "text": "\n".join(current_paras)})
                current_paras = []
            current_heading = txt
        else:
            current_paras.append(txt)
    if current_paras:
        units.append({"section": current_heading, "text": "\n".join(current_paras)})
    return units

def extract_pptx(file_path):
    prs = pptx.Presentation(file_path)
    units = []
    for i, slide in enumerate(prs.slides):
        title = slide.shapes.title.text.strip() if slide.shapes.title and slide.shapes.title.text else f"Slide {i+1}"
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    t = paragraph.text.strip()
                    if t:
                        texts.append(t)
        slide_text = "\n".join(texts)
        if slide_text.strip():
            units.append({"slide": i + 1, "title": title, "text": slide_text})
    return units

def extract_html(file_path):
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        soup = BeautifulSoup(f.read(), "html.parser")
    for s in soup(["script", "style", "meta", "noscript"]):
        s.decompose()
    
    units = []
    current_heading = soup.title.string.strip() if soup.title and soup.title.string else "Overview"
    current_texts = []
    
    for el in soup.find_all(["h1", "h2", "h3", "h4", "p", "li"]):
        tag = el.name.lower()
        txt = el.get_text().strip()
        if not txt:
            continue
        if tag in ["h1", "h2", "h3", "h4"]:
            if current_texts:
                units.append({"section": current_heading, "text": "\n".join(current_texts)})
                current_texts = []
            current_heading = txt
        else:
            current_texts.append(txt)
    if current_texts:
        units.append({"section": current_heading, "text": "\n".join(current_texts)})
    return units

def extract_pdf(file_path):
    reader = PdfReader(file_path)
    units = []
    for i, page in enumerate(reader.pages):
        txt = (page.extract_text() or "").strip()
        if txt:
            units.append({"page": i + 1, "text": txt})
    return units

# Run on sample files
docx_units = extract_docx("DATA/true_data/cronjobs.docx")
docx_chunks = []
for u in docx_units:
    for c in text_splitter.split_text(u["text"]):
        docx_chunks.append({
            "source": "cronjobs.docx",
            "section": u["section"],
            "text": c
        })

pptx_units = extract_pptx("DATA/true_data/architecture.pptx")
pptx_chunks = []
for u in pptx_units:
    for c in text_splitter.split_text(u["text"]):
        pptx_chunks.append({
            "source": "architecture.pptx",
            "slide": u["slide"],
            "slide_title": u["title"],
            "text": c
        })

html_units = extract_html("DATA/true_data/job_management.html")
html_chunks = []
for u in html_units:
    for c in text_splitter.split_text(u["text"]):
        html_chunks.append({
            "source": "job_management.html",
            "section": u["section"],
            "text": c
        })

print(f"Total DOCX chunks: {len(docx_chunks)}")
print(f"Total PPTX chunks: {len(pptx_chunks)}")
print(f"Total HTML chunks: {len(html_chunks)}")

print("\n--- SAMPLE DOCX CHUNK PAYLOAD ---")
print(json.dumps(docx_chunks[2], indent=2))

print("\n--- SAMPLE PPTX CHUNK PAYLOAD ---")
print(json.dumps(pptx_chunks[1], indent=2))

print("\n--- SAMPLE HTML CHUNK PAYLOAD ---")
print(json.dumps(html_chunks[3], indent=2))
